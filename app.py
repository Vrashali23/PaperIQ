import heapq
import re
from collections import Counter

import docx
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import pdfplumber
import plotly.graph_objects as go
import streamlit as st
from fpdf import FPDF
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from textblob import TextBlob
from wordcloud import WordCloud


def generate_research_feedback(results):
    feedback = []
    score = results.get("scores", {}).get("Composite", 0)
    semantic = results.get("semantic_strength", 0)
    citations = results.get("citation_analysis", {}).get("total_citations", 0)
    if score < 70:
        feedback.append("Overall structure needs improvement.")
    if semantic < 60:
        feedback.append("Conceptual depth is limited. Add more technical explanations.")
    if citations < 5:
        feedback.append("Add more scholarly references.")
    if "methodology" not in results.get("sections", {}):
        feedback.append("Methodology section is weak or missing.")
    if not feedback:
        feedback.append("Paper demonstrates strong academic quality.")
    return feedback


def recommend_journal(domain):
    journal_map = {
        "Engineering": "IEEE Transactions",
        "AI": "Elsevier Artificial Intelligence Journal",
        "Medical": "Springer Medical Informatics",
        "Legal": "Journal of Legal Analytics",
    }
    return journal_map.get(domain, "Scopus Indexed Multidisciplinary Journal")


st.set_page_config(page_title="PaperIQ", layout="wide")


if "signed_up" not in st.session_state:
    st.session_state.signed_up = False
if "user_data" not in st.session_state:
    st.session_state.user_data = {}
if "analysis_history" not in st.session_state:
    st.session_state.analysis_history = {}
if not st.session_state.signed_up:
    st.title("📝 PaperIQ Signup")
    username = st.text_input("Username")
    password = st.text_input("Password", type="password")
    role = st.selectbox("Role", ["Student", "Researcher", "Faculty"])
    if st.button("Sign Up", type="primary", key="signup_btn"):
        if username and password:
            st.session_state.user_data = {"username": username, "role": role}
            st.session_state.signed_up = True
            st.success("Signup Successful! Redirecting...")
            st.rerun()
            for key in list(st.session_state.keys()):
                if key not in ["authenticated", "username"]:
                    del st.session_state[key]
        else:
            st.error("Please fill all fields.")
    st.stop()


def extract_text_from_pdf(file):
    text = ""
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    return text


def extract_text_from_docx(file):
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs])


def clean_text(text):
    return re.sub(r"\n+", "\n", text).strip()


def extract_sections(text):
    lines = text.split("\n")
    sections = {}
    current_header = "Preamble"
    current_content = []
    common_headers = [
        "ABSTRACT",
        "INTRODUCTION",
        "LITERATURE REVIEW",
        "METHODOLOGY",
        "RESULTS",
        "DISCUSSION",
        "CONCLUSION",
        "REFERENCES",
        "Abstract",
        "Introduction",
        "Methodology",
        "Conclusion",
    ]
    for line in lines:
        line = line.strip()
        if not line:
            continue
        is_header = False
        if re.match(r"^\d+(\.\d+)*\s+[A-Za-z]", line) and len(line) < 60:
            is_header = True
        elif line in common_headers or (
            line.isupper() and len(line) < 40 and len(line) > 3
        ):
            is_header = True
        if is_header:
            if current_content:
                sections[current_header] = " ".join(current_content)
            current_header = line
            current_content = []
        else:
            current_content.append(line)
    if current_content:
        sections[current_header] = " ".join(current_content)
    return sections


def summarize_text(text, num_sentences=3):
    if not text:
        return "No content to summarize."
    blob = TextBlob(text)
    sentences = blob.sentences
    if len(sentences) <= num_sentences:
        return text
    word_frequencies = {}
    stop_words = set(
        [
            "the",
            "is",
            "in",
            "and",
            "to",
            "of",
            "a",
            "for",
            "on",
            "with",
            "as",
            "by",
            "at",
            "this",
            "that",
            "it",
            "from",
            "an",
            "be",
            "are",
            "was",
        ]
    )
    for word in blob.words:
        word = word.lower()
        if word not in stop_words and word.isalpha():
            word_frequencies[word] = word_frequencies.get(word, 0) + 1
    if not word_frequencies:
        return text
    max_frequency = max(word_frequencies.values())
    for word in word_frequencies:
        word_frequencies[word] = word_frequencies[word] / max_frequency
    sentence_scores = {}
    for sent in sentences:
        for word in sent.words:
            word = word.lower()
            if word in word_frequencies:
                sentence_scores[sent] = (
                    sentence_scores.get(sent, 0) + word_frequencies[word]
                )
    top_sentences = heapq.nlargest(
        num_sentences, sentence_scores, key=sentence_scores.get
    )
    return " ".join([str(s) for s in top_sentences])


def get_important_sentences(text, num_sentences=3):
    sentences = re.split(r"(?<=[.!?]) +", text)
    if len(sentences) <= num_sentences:
        return sentences
    word_freq = {}
    words = re.findall(r"\w+", text.lower())
    for word in words:
        word_freq[word] = word_freq.get(word, 0) + 1
    sentence_scores = {}
    for sentence in sentences:
        for word in re.findall(r"\w+", sentence.lower()):
            if word in word_freq:
                sentence_scores[sentence] = (
                    sentence_scores.get(sentence, 0) + word_freq[word]
                )
    important = heapq.nlargest(num_sentences, sentence_scores, key=sentence_scores.get)
    return important


def count_syllables(word):
    word = word.lower()
    vowels = "aeiou"
    count = 0
    if word[0] in vowels:
        count += 1
    for index in range(1, len(word)):
        if word[index] in vowels and word[index - 1] not in vowels:
            count += 1
    if word.endswith("e"):
        count -= 1
    if count == 0:
        count += 1
    return count


def calculate_readability(text):
    sentences = re.split(r"[.!?]+", text)
    words = text.split()
    total_words = len(words)
    total_sentences = len(sentences)
    syllables = sum(count_syllables(word) for word in words)
    if total_sentences == 0 or total_words == 0:
        return 0
    flesch_score = (
        206.835
        - (1.015 * (total_words / total_sentences))
        - (84.6 * (syllables / total_words))
    )
    return round(flesch_score, 2)


def get_grade(score):
    if score >= 85:
        return "A"
    elif score >= 70:
        return "B"
    elif score >= 55:
        return "C"
    elif score >= 40:
        return "D"
    else:
        return "E"


def analyze_full_document(text):
    blob = TextBlob(text)
    sentences = blob.sentences
    words = blob.words
    word_count = len(words)
    sentence_count = len(sentences)
    if sentence_count == 0:
        return None
    avg_sentence_len = np.mean([len(s.words) for s in sentences])
    avg_word_len = np.mean([len(w) for w in words])
    sentiment = blob.sentiment.polarity
    language_score = min(
        100, (avg_sentence_len * 1.5) + (avg_word_len * 5) + (50 + sentiment * 20)
    )
    transitions = [
        "however",
        "therefore",
        "thus",
        "consequently",
        "furthermore",
        "meanwhile",
    ]
    transition_count = sum(text.lower().count(t) for t in transitions)
    coherence_score = min(100, (transition_count * 4) + (sentence_count * 0.1) + 40)
    reasoning_keywords = [
        "because",
        "since",
        "implies",
        "due to",
        "as a result",
        "evidence",
    ]
    reasoning_count = sum(text.lower().count(k) for k in reasoning_keywords)
    reasoning_score = min(100, (reasoning_count * 6) + 30)
    complex_words = [w for w in words if len(w) > 6]
    lexical_score = (
        min(100, (len(complex_words) / word_count) * 300) if word_count else 0
    )
    readability_score = calculate_readability(text)
    final_score = (
        language_score * 0.3
        + coherence_score * 0.2
        + reasoning_score * 0.2
        + lexical_score * 0.15
        + readability_score * 0.15
    )
    stats = {
        "word_count": word_count,
        "sentence_count": sentence_count,
        "avg_sentence_len": round(avg_sentence_len, 2),
        "avg_word_len": round(avg_word_len, 2),
        "vocab_diversity": round(len(set([w.lower() for w in words])) / word_count, 2)
        if word_count
        else 0,
        "complex_word_ratio": round(len(complex_words) / word_count, 2)
        if word_count
        else 0,
    }
    sections_data = extract_sections(text)
    return {
        "scores": {
            "Language": round(language_score, 2),
            "Coherence": round(coherence_score, 2),
            "Reasoning": round(reasoning_score, 2),
            "Sophistication": round(lexical_score, 2),
            "Readability": round(readability_score, 2),
            "Composite": round(final_score, 2),
        },
        "stats": stats,
        "sentiment": round(sentiment, 2),
        "blob": blob,
        "sections": sections_data,
    }


def generate_full_report(res, filename):
    scores = res["scores"]
    stats = res["stats"]
    keywords = res.get("keywords", [])
    domain = res.get("domain", "Unknown")
    publisher = res.get("publisher", "Not detected")
    sections = res.get("sections", {})
    report = f"""
PAPER ANALYSIS REPORT
========================
File Name: {filename}
Publisher: {publisher}
Domain: {domain}
----------------------------------------
DOCUMENT STATISTICS
----------------------------------------
Word Count: {stats["word_count"]}
Sentence Count: {stats["sentence_count"]}
----------------------------------------
SCORES
----------------------------------------
Language: {scores["Language"]}
Coherence: {scores["Coherence"]}
Reasoning: {scores["Reasoning"]}
Sophistication: {scores["Sophistication"]}
Readability: {scores["Readability"]}
Composite: {scores["Composite"]}
----------------------------------------
KEYWORDS
----------------------------------------
{", ".join(keywords)}
----------------------------------------
SECTION SUMMARIES
----------------------------------------
"""
    for title, summary in sections.items():
        report += f"\n\n{title}\n{'-' * len(title)}\n{summary}\n"
    report += "\n\n----------------------------------------\nEND OF REPORT\n----------------------------------------"
    return report


def clean_for_pdf(text):
    if not text:
        return ""
    return text.encode("latin-1", "ignore").decode("latin-1")


def generate_pdf_report(res, filename):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=10)
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "PaperIQ Analysis Report", ln=True)
    pdf.set_font("Arial", size=12)
    pdf.ln(5)
    pdf.multi_cell(0, 8, f"File Name: {filename}")
    pdf.multi_cell(0, 8, clean_for_pdf(f"Title: {res.get('title')}"))
    pdf.multi_cell(0, 8, f"Domain: {res.get('domain')}")
    pdf.multi_cell(0, 8, f"Publisher: {res.get('publisher')}")
    pdf.ln(5)
    pdf.cell(0, 10, "Scores:", ln=True)
    for k, v in res["scores"].items():
        pdf.multi_cell(0, 8, f"{k}: {v}")
    pdf.ln(5)
    pdf.ln(10)
    pdf.cell(0, 10, "Section Summaries:", ln=True)
    for title, summary in res["sections"].items():
        pdf.set_font("Arial", "B", 12)
        pdf.multi_cell(0, 8, clean_for_pdf(title))
        pdf.set_font("Arial", size=11)
        pdf.multi_cell(0, 8, clean_for_pdf(summary))
        pdf.ln(3)
    pdf_output = pdf.output(dest="S")
    return bytes(pdf_output, "latin-1")


def extract_keywords_and_domain(text, top_n=15):
    blob = TextBlob(text)
    stop_words = set(
        [
            "the",
            "is",
            "in",
            "and",
            "to",
            "of",
            "a",
            "for",
            "on",
            "with",
            "as",
            "by",
            "at",
            "this",
            "that",
            "it",
            "from",
            "an",
            "be",
            "are",
            "was",
        ]
    )
    word_freq = {}
    for word in blob.words:
        word = word.lower()
        if word.isalpha() and word not in stop_words and len(word) > 3:
            word_freq[word] = word_freq.get(word, 0) + 1
    sorted_words = sorted(word_freq.items(), key=lambda x: x[1], reverse=True)
    keywords = [w[0] for w in sorted_words[:top_n]]
    domains = {
        "Computer Science": ["algorithm", "data", "model", "network", "ai", "learning"],
        "Healthcare": ["patient", "medical", "treatment", "disease", "clinical"],
        "Finance": ["market", "investment", "economic", "revenue", "financial"],
        "Education": ["learning", "students", "teaching", "curriculum"],
        "Engineering": ["system", "design", "performance", "analysis"],
    }
    detected_domain = "General"
    for domain, terms in domains.items():
        for term in terms:
            if term in text.lower():
                detected_domain = domain
                break
    return keywords, detected_domain


def reviewer_comments(score):
    if score > 80:
        return "Strong paper. Minor revisions recommended."
    elif score > 60:
        return "Good work, but clarity and structure need improvement."
    else:
        return "Major revisions required before acceptance."


def semantic_answer(question, full_text, sections):
    question = question.lower()
    section_keywords = {
        "abstract": "Abstract",
        "introduction": "Introduction",
        "method": "Methodology",
        "methods": "Methodology",
        "methodology": "Methodology",
        "results": "Results",
        "discussion": "Discussion",
        "conclusion": "Conclusion",
    }
    for key, section_name in section_keywords.items():
        if key in question:
            for title, content in sections.items():
                if section_name.lower() in title.lower():
                    return content[:1000]
    sentences = re.split(r"(?<=[.!?]) +", full_text)
    vectorizer = TfidfVectorizer().fit_transform(sentences + [question])
    vectors = vectorizer.toarray()
    cosine_sim = cosine_similarity([vectors[-1]], vectors[:-1])[0]
    best_match_index = np.argmax(cosine_sim)
    return sentences[best_match_index]


STOPWORDS = set(
    [
        "the",
        "is",
        "in",
        "and",
        "to",
        "of",
        "a",
        "an",
        "by",
        "this",
        "that",
        "it",
        "as",
        "are",
        "was",
    ]
)


def detect_repetition(text, threshold=5):
    words = re.findall(r"\b[a-zA-Z]{4,}\b", text.lower())
    words = [w for w in words if w not in STOPWORDS]
    word_counts = Counter(words)
    repeated = {
        word: count for word, count in word_counts.items() if count >= threshold
    }
    return repeated


def generate_structured_suggestions(text):
    suggestions = []
    sentences = re.split(r"[.!?]", text)
    long_sentences = [s for s in sentences if len(s.split()) > 30]
    if long_sentences:
        suggestions.append(
            "• Consider breaking long sentences into shorter ones for clarity."
        )
    repeated_words = detect_repetition(text)
    if repeated_words:
        suggestions.append(
            "• Reduce repetition of words like: "
            + ", ".join(list(repeated_words.keys())[:5])
        )
    if " was " in text or " were " in text:
        suggestions.append("• Review passive voice usage for stronger academic tone.")
    if not suggestions:
        suggestions.append(
            "• Writing structure looks good. Minor refinements may improve clarity."
        )
    return suggestions


def detect_publisher(text):
    publishers = {
        "IEEE": ["ieee", "ieeexplore"],
        "Springer": ["springer"],
        "Elsevier": ["elsevier", "sciencedirect"],
        "Wiley": ["wiley"],
        "Taylor & Francis": ["taylor", "francis"],
        "ACM": ["acm", "association for computing machinery"],
        "Nature": ["nature publishing"],
        "MDPI": ["mdpi"],
        "Oxford University Press": ["oxford university press"],
    }
    text_lower = text.lower()
    for publisher, keywords in publishers.items():
        for word in keywords:
            if word in text_lower:
                return publisher
    return "Not Detected"


def section_wise_sentiment(sections):
    sentiment_results = {}
    blob = TextBlob(full_text)
    overall_sentiment = blob.sentiment.polarity
    return sentiment_results


def novelty_score(text):
    words = re.findall(r"\b\w+\b", text.lower())
    if not words:
        return 0
    return round((len(set(words)) / len(words)) * 100, 2)


def analyze_citations(text):
    citation_patterns = [r"\(\d{4}\)", r"\[\d+\]", r"\(\w+ et al\., \d{4}\)"]
    citation_count = 0
    for pattern in citation_patterns:
        citation_count += len(re.findall(pattern, text))
    total_words = len(text.split())
    citation_density = round(citation_count / total_words, 4) if total_words > 0 else 0
    impact_score = min(100, citation_count * 2)
    years = re.findall(r"(19\d{2}|20\d{2})", text)
    if years:
        avg_year = sum(map(int, years)) / len(years)
    else:
        avg_year = 0
    return {
        "total_citations": citation_count,
        "citation_density": citation_density,
        "impact_score": impact_score,
        "average_year": round(avg_year, 1),
    }


def calculate_semantic_strength(text):
    blob = TextBlob(text)
    nouns = [word.lower() for word, tag in blob.tags if tag.startswith("NN")]
    unique_nouns = len(set(nouns))
    total_nouns = len(nouns)
    if total_nouns == 0:
        return 0
    return round(min(100, (unique_nouns / total_nouns) * 150), 2)


def extract_topic_focus(text):
    keywords, _ = extract_keywords_and_domain(text, top_n=15)
    return keywords


with st.sidebar:
    st.markdown("User Dashboard")
    st.write(f"👤 {st.session_state.user_data['username']}")
    st.write(f"🎓 Role: {st.session_state.user_data['role']}")

    if st.sidebar.button("Logout"):
        for key in list(st.session_state.keys()):
            del st.session_state[key]
        st.rerun()
    st.divider()

if "analysis_history" not in st.session_state:
    st.session_state.analysis_history = {}
if "comparison_results" not in st.session_state:
    st.session_state.comparison_results = []
st.markdown("### 📤 Upload Research Paper")
uploaded_files = st.file_uploader("", type=["pdf"], accept_multiple_files=True)
if uploaded_files:
    if st.button("Analyze Document", type="primary", key="analyze_btn"):
        with st.spinner("Analyzing document..."):
            username = st.session_state.user_data["username"]
            if username not in st.session_state.analysis_history:
                st.session_state.analysis_history[username] = []
            st.session_state.comparison_results = []
            for file in uploaded_files:
                if file.name.endswith(".pdf"):
                    text = extract_text_from_pdf(file)
                    cleaned_text = clean_text(text)
                    results = analyze_full_document(cleaned_text)
                    if results:
                        keywords, domain = extract_keywords_and_domain(cleaned_text)
                        publisher = detect_publisher(cleaned_text)
                        results["keywords"] = keywords
                        results["domain"] = domain
                        results["publisher"] = publisher
                        results["full_text"] = cleaned_text
                        results["citation_analysis"] = analyze_citations(cleaned_text)
                        results["semantic_strength"] = calculate_semantic_strength(
                            cleaned_text
                        )
                        results["topic_focus"] = extract_topic_focus(cleaned_text)
                        results["ai_feedback"] = generate_research_feedback(results)
                        results["recommended_journal"] = recommend_journal(
                            results["domain"]
                        )
                        results["novelty_score"] = novelty_score(cleaned_text)
                        st.session_state.comparison_results.append(results)
                        st.session_state["results"] = results
                        st.session_state["filename"] = file.name
                        pdf_bytes = generate_pdf_report(results, file.name)
                        st.session_state.analysis_history[username].append(
                            {
                                "filename": file.name,
                                "domain": domain,
                                "score": results["scores"]["Composite"],
                                "pdf": pdf_bytes,
                            }
                        )
                else:
                    st.error(f"Unable to analyze file: {file.name}")


if "results" in st.session_state:
    res = st.session_state["results"]
    scores = res["scores"]
    stats = res["stats"]

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Language", f"{scores['Language']}/100")
    col2.metric("Coherence", f"{scores['Coherence']}/100")
    col3.metric("Reasoning", f"{scores['Reasoning']}/100")
    col4.metric("Sophistication", f"{scores['Sophistication']}/100")
    if st.button("Generate PDF Report"):
        pdf_report = generate_pdf_report(res, st.session_state["filename"])
        st.download_button(
            "Download Full Report",
            data=pdf_report,
            file_name="PaperIQ_Report.pdf",
            mime="application/pdf",
        )
if "results" not in st.session_state:
    st.info("Upload and analyze a paper to unlock features.")
    st.stop()


# ---------------- NAVIGATION ----------------
st.sidebar.markdown("## Navigation")

if "page" not in st.session_state:
    st.session_state.page = "Overview"

# 🔹 Main Sections (Reduced to 5)
if st.sidebar.button("Overview", width="stretch"):
    st.session_state.page = "Overview"

if st.sidebar.button("Analysis", width="stretch"):
    st.session_state.page = "Analysis"

if st.sidebar.button("Content Insights", width="stretch"):
    st.session_state.page = "Content"

if st.sidebar.button("Chatbot", width="stretch"):
    st.session_state.page = "AI"

if st.sidebar.button("Workspace", width="stretch"):
    st.session_state.page = "Workspace"


# ---------------- LOAD RESULTS ----------------
if "results" in st.session_state:
    res = st.session_state["results"]
    scores = res["scores"]
    stats = res["stats"]


# ---------------- PAGE RENDERING ----------------

# 📄 OVERVIEW
if st.session_state.page == "Overview":
    st.header("📄 Overview")

    # ---------------- SCORE BREAKDOWN ----------------
    st.markdown("### 📊 Score Breakdown")

    col1, col2, col3, col4 = st.columns(4)

    col1.metric("Language", f"{scores['Language']}")
    col2.metric("Coherence", f"{scores['Coherence']}")
    col3.metric("Reasoning", f"{scores['Reasoning']}")
    col4.metric("Sophistication", f"{scores['Sophistication']}")

    st.markdown("")

    col5, _ = st.columns([1, 3])
    col5.metric("Readability", f"{scores['Readability']}")

    st.divider()

    # ---------------- METADATA ----------------
    st.markdown("### 📄 Paper Details")

    # 🔍 Publisher detection (simple heuristic)
    text_sample = res.get("full_text", "")[:1000].lower()

    publisher = "Unknown"

    publisher = detect_publisher(res.get("full_text", ""))
    st.write(f"**Publisher:** {publisher}")
    st.write(f"**Domain:** {res['domain']}")
    st.write(f"**Word Count:** {stats['word_count']}")
    st.write(f"**Sentence Count:** {stats['sentence_count']}")

    st.divider()

    # ---------------- WORD CLOUD ----------------
    st.markdown("### ☁️ Key Concepts")

    text = res.get("full_text", "")

    if not text or len(text.strip()) == 0:
        st.warning("No text available to generate word cloud.")

    else:
        # 🔥 Reduce clutter: limit text + filter small words
        words = text.split()
        filtered_words = [w for w in words if len(w) > 4]

        reduced_text = " ".join(filtered_words[:1000])  # limit words

        try:
            wordcloud = WordCloud(
                width=800,
                height=400,
                background_color="white",
                max_words=50,  # 🔥 LIMIT WORDS
                colormap="viridis",
            ).generate(reduced_text)

            fig, ax = plt.subplots(figsize=(10, 5))
            ax.imshow(wordcloud, interpolation="bilinear")
            ax.axis("off")

            st.pyplot(fig)

        except Exception as e:
            st.error(f"WordCloud error: {e}")

# 📊 ANALYSIS
elif st.session_state.page == "Analysis":
    st.header("📊 Analysis")

    # ---------------- TOP METRICS ----------------
    st.markdown("### 📈 Overall Performance")

    grade = get_grade(scores["Composite"])

    col1, col2, col3 = st.columns(3)

    col1.metric("Composite Score", f"{scores['Composite']}/100")
    col2.metric("Grade", grade)
    col3.metric("Sentiment", res["sentiment"])

    # optional feedback
    if grade == "A":
        st.success("Excellent paper quality")
    elif grade == "B":
        st.info("Good paper with minor improvements")
    else:
        st.warning("Needs improvement")

    st.divider()

    # ---------------- VISUALS (SIDE BY SIDE) ----------------
    st.markdown("### 📊 Performance Overview")

    col1, col2 = st.columns([2, 1])  # radar bigger, gauge smaller

    # 🔵 Radar Chart
    with col1:
        categories = [
            "Language",
            "Coherence",
            "Reasoning",
            "Sophistication",
            "Readability",
        ]

        values = [scores[c] for c in categories]

        categories_closed = categories + [categories[0]]
        values_closed = values + [values[0]]

        fig = go.Figure()

        fig.add_trace(
            go.Scatterpolar(
                r=values_closed,
                theta=categories_closed,
                fill="toself",
                name="Your Paper",
            )
        )

        ideal = [85, 85, 85, 85, 85]
        ideal_closed = ideal + [ideal[0]]

        fig.add_trace(
            go.Scatterpolar(
                r=ideal_closed,
                theta=categories_closed,
                line=dict(dash="dash"),
                name="Ideal",
            )
        )

        fig.update_layout(
            polar=dict(radialaxis=dict(visible=True, range=[0, 100])), height=420
        )

        st.plotly_chart(fig, use_container_width=True)

    # 🟣 Gauge Chart
    with col2:
        fig = go.Figure(
            go.Indicator(
                mode="gauge+number",
                value=scores["Composite"],
                title={"text": "Quality"},
                gauge={
                    "axis": {"range": [0, 100]},
                    "bar": {"color": "darkblue"},
                },
            )
        )

        fig.update_layout(height=420)

        st.plotly_chart(fig, use_container_width=True)

    st.divider()

    # ---------------- SCORE BREAKDOWN ----------------
    st.markdown("### 📊 Detailed Scores")

    col1, col2, col3, col4 = st.columns(4)

    col1.metric("Language", scores["Language"])
    col2.metric("Coherence", scores["Coherence"])
    col3.metric("Reasoning", scores["Reasoning"])
    col4.metric("Sophistication", scores["Sophistication"])

    st.markdown("")

    col5, _ = st.columns([1, 3])
    col5.metric("Readability", scores["Readability"])

    st.divider()

    # ---------------- INSIGHTS ----------------
    st.markdown("### 🧠 Insights")

    col1, col2 = st.columns(2)

    col1.metric("Semantic Strength", f"{res.get('semantic_strength', 0)}/100")
    col2.metric("Sentence Count", stats["sentence_count"])


# 📑 CONTENT INSIGHTS
elif st.session_state.page == "Content":
    st.header("📑 Content Insights")

    tab1, tab2 = st.tabs(["📑 Section Summaries", "💡 Suggestions"])

    # 📑 Summaries Tab
    with tab1:
        for title, content in res["sections"].items():
            summary = summarize_text(content, 5)
            st.write(f"**{title}**")
            st.write(summary)
            st.divider()

    # 💡 Suggestions Tab
    with tab2:
        for point in res["ai_feedback"]:
            st.write("•", point)

# 🤖 AI TOOLS
elif st.session_state.page == "AI":
    st.header("🤖 AI Tools")

    user_question = st.text_input("Ask a question about the paper")

    if user_question:
        answer = semantic_answer(user_question, res["full_text"], res["sections"])
        st.write(answer)


# 📂 WORKSPACE
elif st.session_state.page == "Workspace":
    st.header("📂 Workspace")

    with st.expander("📂 History", expanded=True):
        username = st.session_state.user_data["username"]
        history = st.session_state.analysis_history.get(username, [])

        if not history:
            st.info("No history available")
        else:
            for item in history:
                st.write(item["filename"], "-", item["score"])

    with st.expander("📊 Comparison"):
        st.write("Comparison section here")
